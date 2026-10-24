"""Build the architecture deck and the jury presentation from saved metrics."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from src.config import DOCS_DIR, MODEL_DIR

TEAL = RGBColor(0x0F, 0x6E, 0x6E)
INK = RGBColor(0x10, 0x21, 0x24)
MUTED = RGBColor(0x4D, 0x61, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SAND = RGBColor(0xF4, 0xF7, 0xF8)
AMBER = RGBColor(0xB5, 0x47, 0x08)
CARD = RGBColor(0xE6, 0xF2, 0xF2)


def _deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "HealthTwin"
    prs.core_properties.author = "Ritik Awachat"
    prs.core_properties.subject = "Digital Twin Challenge 2026"
    return prs


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _text(slide, text, x, y, w, h, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    try:
        frame._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr"}[anchor])
    except Exception:
        pass
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def _rect(slide, x, y, w, h, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(shape, color)
    shape.adjustments[0] = 0.08
    return shape


def _title_bar(slide, kicker: str, title: str) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.15))
    _fill(bar, TEAL)
    _text(slide, kicker, 0.45, 0.12, 12, 0.3, size=12, bold=True, color=RGBColor(0xC9, 0xEB, 0xE6))
    _text(slide, title, 0.45, 0.38, 12, 0.6, size=28, bold=True, color=WHITE)


def _bullets(slide, items: list[str], x, y, w, h, size=18) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.space_after = Pt(16)
        run = p.add_run()
        run.text = item
        run.font.size = Pt(size)
        run.font.color.rgb = INK
        run.font.name = "Calibri"


def _load_metrics() -> dict:
    path = MODEL_DIR / "metrics.json"
    if not path.exists():
        raise FileNotFoundError("Train the model first: python -m src.train")
    return json.loads(path.read_text(encoding="utf-8"))


def _fmt(metrics: dict, *path, digits=3) -> str:
    value = metrics
    for key in path:
        value = value[key]
    if value is None:
        return "n/a"
    if isinstance(value, float):
        return f"{value:.{digits}f}"
    return str(value)


def _flow(slide, top: float) -> None:
    boxes = [
        (0.4, "Clinic chart", "Age, labs, medicines\nBlood pressure\nFamily history\nSimulated gene flags"),
        (3.55, "Wearable stream", "Glucose every 15 min\nHeart rate, steps, sleep\nPast only\nNothing from later"),
        (6.7, "Model", "Patient-level split\nSpike at 180 or above\nDrop under 70\nPeak sketch"),
        (9.85, "Dashboard", "Who should get a message?\nThen one person\nGreen line so far\nOrange estimate"),
    ]
    for x, title, body in boxes:
        _rect(slide, x, top, 2.9, 3.15, WHITE)
        _text(slide, title, x + 0.18, top + 0.15, 2.55, 0.45, size=20, bold=True, color=TEAL)
        _text(slide, body, x + 0.18, top + 0.7, 2.55, 2.2, size=18, color=INK)
    for x in (3.15, 6.3, 9.45):
        _text(slide, "→", x, top + 1.2, 0.45, 0.5, size=28, bold=True, color=TEAL, align=PP_ALIGN.CENTER)


def build_architecture(metrics: dict, path: Path) -> None:
    prs = _deck()
    slide = _blank(prs)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    _fill(bg, SAND)
    _title_bar(slide, "HEALTHTWIN  ·  ARCHITECTURE", "Two streams in. A two-hour warning out.")
    _flow(slide, 1.7)
    _text(
        slide,
        "Synthetic data only. Features stop at the current minute.",
        0.5,
        5.15,
        12.3,
        0.4,
        size=18,
        color=MUTED,
    )
    _text(
        slide,
        "Two adverse events: glucose ≥180 (spike) and <70 (low).",
        0.5,
        5.6,
        12.3,
        0.4,
        size=20,
        bold=True,
        color=INK,
    )
    _text(
        slide,
        f"Held-out ROC-AUC  fusion {_fmt(metrics, 'fusion_test', 'roc_auc')}   ·   "
        f"glucose sensor only {_fmt(metrics, 'cgm_only_test', 'roc_auc')}   ·   "
        f"still under 140: fusion {_fmt(metrics, 'lead_time_auc_fusion')}",
        0.5,
        6.35,
        12.3,
        0.5,
        size=18,
        bold=True,
        color=INK,
    )
    prs.save(path)


def build_presentation(metrics: dict, path: Path) -> None:
    prs = _deck()
    fusion = metrics["fusion_test"]

    slide = _blank(prs)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    _fill(bg, TEAL)
    _text(slide, "HAPPIEST HEALTH  ·  DIGITAL TWIN CHALLENGE 2026", 0.7, 1.5, 12, 0.4, size=16, bold=True, color=RGBColor(0xC9, 0xEB, 0xE6))
    _text(slide, "HealthTwin", 0.7, 2.0, 12, 1.0, size=54, bold=True, color=WHITE)
    _text(
        slide,
        "See the change. Understand the risk. Explore what comes next.",
        0.7,
        3.2,
        11,
        0.6,
        size=26,
        color=WHITE,
    )
    _text(
        slide,
        "A two-hour blood-sugar warning for one doctor and one person.",
        0.7,
        4.05,
        11,
        0.5,
        size=22,
        color=WHITE,
    )
    _text(
        slide,
        "Team Ritik Awachat   ·   Vijay Patil School of Management   ·   solo",
        0.7,
        6.2,
        11.5,
        0.4,
        size=20,
        color=RGBColor(0xC9, 0xEB, 0xE6),
    )

    slide = _blank(prs)
    _title_bar(slide, "01  ·  PROBLEM", "The average does not say what the next two hours will do.")
    _bullets(
        slide,
        [
            "HbA1c is a three-month average. It does not say the next meal will pass 180.",
            "HealthTwin joins the clinic chart to a wearable stream.",
            "The warning is two hours ahead: long enough to walk, eat later, or send a message.",
            "The screen is for the doctor. It is not a diagnosis.",
        ],
        0.7,
        1.8,
        12,
        4.5,
        size=26,
    )

    slide = _blank(prs)
    _title_bar(slide, "02  ·  USE CASE", "One condition. One person. One decision.")
    _bullets(
        slide,
        [
            "A working-age adult with Type 2 diabetes, a glucose sensor, and a watch.",
            "Sugar can still look ordinary when a spike is already likely.",
            "The doctor sees who to message, then the chance, the peak, and the reasons.",
            "It does not change a dose, and it does not forecast a heart attack.",
        ],
        0.7,
        1.8,
        12,
        4.5,
        size=26,
    )

    slide = _blank(prs)
    _title_bar(slide, "03  ·  SANDBOX", "No real patient record is in this prototype.")
    _bullets(
        slide,
        [
            "Real charts stay out. The clinic file and the wearable stream are synthetic.",
            "Chart: labs, medicines, blood pressure, family history, simulated gene flags.",
            "Wearable: glucose, heart rate, steps, and sleep, every 15 minutes.",
            "Meals change the sugar. The model is not given a column called meal.",
        ],
        0.7,
        1.8,
        12,
        4.5,
        size=26,
    )

    slide = _blank(prs)
    _title_bar(slide, "04  ·  ARCHITECTURE", "Two streams, then the model, then the screen.")
    _flow(slide, 1.55)
    _text(
        slide,
        "Features stop at the current minute. The label is the next two hours.",
        0.55,
        5.0,
        12,
        0.4,
        size=20,
        color=MUTED,
    )
    _text(
        slide,
        "Two adverse events: glucose ≥180 (spike) and <70 (low).",
        0.55,
        5.55,
        12,
        0.45,
        size=24,
        bold=True,
        color=INK,
    )

    slide = _blank(prs)
    _title_bar(slide, "05  ·  FUSION", "The join is the product. Either stream alone is not.")
    _bullets(
        slide,
        [
            "Inputs are the chart plus sensors up to this minute. Nothing later is an input.",
            "Two adverse events: glucose ≥180 (spike) and <70 (low).",
            "People are split 70 / 15 / 15. One person is never in two sets.",
            "The alert line is chosen on the validation people, then frozen.",
        ],
        0.7,
        1.8,
        12,
        4.5,
        size=26,
    )

    slide = _blank(prs)
    _title_bar(slide, "06  ·  MODEL", "Three models. One laptop. A frozen threshold.")
    _bullets(
        slide,
        [
            "Gradient boosting calls a spike, calls a drop under 70, and sketches the peak.",
            "A glucose-only model is the baseline, with its own threshold.",
            "The screen explains the call in sentences, not a table of weights.",
            f"Python, pandas, scikit-learn, Streamlit. Seed {metrics['seed']}. {metrics['n_patients']} people, {metrics['n_days']} days, {metrics['n_test_patients']} held out.",
        ],
        0.7,
        1.8,
        12,
        4.8,
        size=24,
    )

    slide = _blank(prs)
    _title_bar(slide, "07  ·  RESULTS", "23 people the model had not seen.")
    cards = [
        (0.45, "Fusion ROC-AUC", _fmt(metrics, "fusion_test", "roc_auc"), "Full twin"),
        (3.6, "Sensor-only ROC-AUC", _fmt(metrics, "cgm_only_test", "roc_auc"), "Glucose features only"),
        (6.75, "In-range fusion", _fmt(metrics, "lead_time_auc_fusion"), "Sugar still under 140"),
        (9.9, "Peak error", f"{metrics['peak_mae']:.1f}", "mg/dL, average"),
    ]
    for x, label, value, hint in cards:
        _rect(slide, x, 1.55, 2.95, 2.35, CARD)
        _text(slide, label, x + 0.15, 1.7, 2.65, 0.45, size=16, color=TEAL, bold=True)
        _text(slide, value, x + 0.15, 2.2, 2.65, 0.75, size=40, bold=True, color=INK)
        _text(slide, hint, x + 0.15, 3.15, 2.65, 0.45, size=16, color=MUTED)
    _bullets(
        slide,
        [
            f"Threshold {metrics['threshold']:.2f}, set on validation: precision {fusion['precision']:.2f}, recall {fusion['recall']:.2f}.",
            f"While sugar is still under 140, fusion is {_fmt(metrics, 'lead_time_auc_fusion')} and the sensor-only model is {_fmt(metrics, 'lead_time_auc_cgm_only')}.",
        ],
        0.6,
        4.25,
        12.1,
        2.4,
        size=22,
    )

    slide = _blank(prs)
    _title_bar(slide, "08  ·  DASHBOARD", "Who should get a message, then one person.")
    lines = [
        "Message them, already too high, or no message. Reyansh Joshi is first.",
        "Sugar 117, a 96% chance it passes 180. A 30-minute walk moves that to 87%.",
        "Green is sugar so far. Orange is the estimate. The recorded line does not move.",
    ]
    for i, line in enumerate(lines):
        _text(slide, line, 0.5, 1.28 + i * 0.38, 12.3, 0.38, size=20, color=INK)
    list_shot = DOCS_DIR / "screen-clinic-list.png"
    walk_shot = DOCS_DIR / "screen-reyansh-walk.png"
    if not list_shot.exists() or not walk_shot.exists():
        raise FileNotFoundError("Dashboard screenshots are missing from docs/.")
    slide.shapes.add_picture(str(list_shot), Inches(0.35), Inches(2.55), Inches(6.2))
    slide.shapes.add_picture(str(walk_shot), Inches(6.75), Inches(2.55), Inches(6.2))

    slide = _blank(prs)
    _title_bar(slide, "09  ·  LIMITS", "Said here, before anyone has to ask.")
    _bullets(
        slide,
        [
            "The score is on a generator we wrote. It is not a ward validation.",
            "Meals follow a clock. A next version needs a meal log.",
            "The gene flags are simulated. The chart is a sketch, not a body.",
            "Do not use this to change a dose.",
        ],
        0.7,
        1.8,
        12,
        4.5,
        size=26,
    )

    slide = _blank(prs)
    _title_bar(slide, "10  ·  CLOSE", "A quiet reading that will not stay quiet.")
    _bullets(
        slide,
        [
            "Two streams. Two warnings. One screen a doctor can stand in front of.",
            "Team name: Ritik Awachat. College: Vijay Patil School of Management. Solo.",
            "The code, the licence, and this deck are in the public repository.",
        ],
        0.7,
        2.0,
        12,
        3.2,
        size=26,
    )
    _text(slide, "Practice data. Not real patients. Not medical advice.", 0.7, 6.4, 12, 0.4, size=20, color=MUTED)

    prs.save(path)


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    metrics = _load_metrics()
    architecture = DOCS_DIR / "HealthTwin_Architecture.pptx"
    presentation = DOCS_DIR / "HealthTwin_Presentation.pptx"
    build_architecture(metrics, architecture)
    build_presentation(metrics, presentation)
    print(f"Wrote {architecture}")
    print(f"Wrote {presentation}")


if __name__ == "__main__":
    main()
